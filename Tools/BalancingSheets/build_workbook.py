from __future__ import annotations

import argparse
import csv
import json
import math
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


TOOL_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = TOOL_DIR.parents[1]
DEFAULT_CATALOG = TOOL_DIR / "sheet_catalog.json"
DEFAULT_OUTPUT = TOOL_DIR / "output" / "NeoSurvive_Balancing.xlsx"

TEXT_COLUMNS = {
    "id",
    "enemyid",
    "weaponid",
    "mapid",
    "name",
    "type",
    "sortinglayername",
    "effecttype",
    "hackedeffecttype",
    "notes",
}
BOOLEAN_COLUMNS = {
    "siegefreezeoutershooters",
    "autofit",
    "droprandomone",
    "dropboth",
    "enablerandomspark",
    "isactive",
    "canhack",
    "blink",
}
TRUE_VALUES = {"true", "1"}
FALSE_VALUES = {"false", "0"}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
INDEX_TITLE_FILL = PatternFill("solid", fgColor="17365D")
INDEX_SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
ERROR_FILL = PatternFill("solid", fgColor="F4CCCC")
DUPLICATE_FILL = PatternFill("solid", fgColor="FCE5CD")


@dataclass
class SheetResult:
    sheet: str
    category: str
    path: str
    row_count: int
    issues: list[str] = field(default_factory=list)

    @property
    def status(self) -> str:
        return "검토 필요" if self.issues else "정상"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="NeoSurvive CSV 29개를 Google Sheets 업로드용 Excel 워크북으로 변환합니다."
    )
    parser.add_argument("--catalog", type=Path, default=DEFAULT_CATALOG)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--verify-only",
        action="store_true",
        help="기존 워크북의 탭 수와 기본 구조만 검증합니다.",
    )
    return parser.parse_args()


def load_catalog(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as file:
        catalog = json.load(file)

    entries = catalog.get("entries")
    if not isinstance(entries, list) or len(entries) != 29:
        raise ValueError(f"카탈로그에는 정확히 29개 항목이 필요합니다: {len(entries or [])}개")

    sheets = [entry["sheet"] for entry in entries]
    paths = [entry["path"] for entry in entries]
    if len(set(sheets)) != len(sheets):
        raise ValueError("카탈로그에 중복된 시트 이름이 있습니다.")
    if len(set(paths)) != len(paths):
        raise ValueError("카탈로그에 중복된 CSV 경로가 있습니다.")
    if any(len(sheet) > 31 for sheet in sheets):
        raise ValueError("Excel 시트 이름은 31자를 넘을 수 없습니다.")

    return catalog


def read_csv(path: Path) -> tuple[list[str], list[list[str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        rows = list(csv.reader(file))

    if not rows:
        raise ValueError(f"비어 있는 CSV입니다: {path}")

    return rows[0], rows[1:]


def parse_numeric(value: str) -> int | float | None:
    text = value.strip()
    if not text:
        return None

    try:
        number = float(text)
    except ValueError:
        return None

    if not math.isfinite(number):
        return None
    if number.is_integer() and "." not in text and "e" not in text.lower():
        return int(number)
    return number


def parse_boolean(value: str) -> bool | None:
    normalized = value.strip().lower()
    if normalized in TRUE_VALUES:
        return True
    if normalized in FALSE_VALUES:
        return False
    return None


def convert_cell(header: str, raw_value: str) -> tuple[Any, str | None]:
    normalized_header = header.strip().lower()

    if normalized_header in TEXT_COLUMNS or not normalized_header:
        return raw_value, None

    if normalized_header in BOOLEAN_COLUMNS:
        value = parse_boolean(raw_value)
        if value is None and raw_value.strip():
            return raw_value, "불리언 값은 TRUE/FALSE 또는 1/0이어야 합니다."
        return value, None

    value = parse_numeric(raw_value)
    if value is None and raw_value.strip():
        return raw_value, "숫자 열에 숫자가 아닌 값이 있습니다."
    return value, None


def key_tuple(headers: list[str], row: list[str], keys: list[str]) -> tuple[str, ...]:
    header_lookup = {header.lower(): index for index, header in enumerate(headers)}
    values: list[str] = []
    for key in keys:
        index = header_lookup[key.lower()]
        values.append(row[index].strip().lower() if index < len(row) else "")
    return tuple(values)


def add_data_validation(worksheet: Any, headers: list[str], max_row: int) -> None:
    validation_end = max(1000, max_row + 200)
    for column_index, header in enumerate(headers, start=1):
        normalized_header = header.strip().lower()
        if not normalized_header:
            continue

        column_letter = get_column_letter(column_index)
        cell_range = f"{column_letter}2:{column_letter}{validation_end}"
        if normalized_header in BOOLEAN_COLUMNS:
            validation = DataValidation(
                type="list",
                formula1='"TRUE,FALSE"',
                allow_blank=True,
                error="TRUE 또는 FALSE만 입력할 수 있습니다.",
                errorTitle="잘못된 불리언 값",
            )
        elif normalized_header not in TEXT_COLUMNS:
            validation = DataValidation(
                type="decimal",
                operator="between",
                formula1="-1000000000",
                formula2="1000000000",
                allow_blank=True,
                error="숫자만 입력할 수 있습니다.",
                errorTitle="잘못된 숫자 값",
            )
        else:
            continue

        validation.errorStyle = "stop"
        validation.showErrorMessage = True
        worksheet.add_data_validation(validation)
        validation.add(cell_range)


def add_duplicate_rule(worksheet: Any, headers: list[str], keys: list[str], max_row: int) -> None:
    if max_row < 2:
        return

    header_lookup = {header.lower(): index + 1 for index, header in enumerate(headers)}
    key_columns = [header_lookup[key.lower()] for key in keys]
    criteria: list[str] = []
    for column in key_columns:
        letter = get_column_letter(column)
        criteria.extend([f"${letter}$2:${letter}${max_row}", f"${letter}2"])

    formula = f"COUNTIFS({','.join(criteria)})>1"
    worksheet.conditional_formatting.add(
        f"A2:{get_column_letter(len(headers))}{max_row}",
        FormulaRule(formula=[formula], fill=DUPLICATE_FILL),
    )


def set_column_widths(worksheet: Any, headers: list[str], rows: list[list[str]]) -> None:
    for column_index, header in enumerate(headers, start=1):
        values = [header]
        values.extend(row[column_index - 1] for row in rows if column_index - 1 < len(row))
        longest = max((len(str(value)) for value in values), default=8)
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(longest + 2, 10), 42
        )


def build_data_sheet(
    workbook: Workbook,
    entry: dict[str, Any],
    headers: list[str],
    rows: list[list[str]],
) -> SheetResult:
    worksheet = workbook.create_sheet(entry["sheet"])
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = True
    worksheet.append(headers)

    result = SheetResult(
        sheet=entry["sheet"],
        category=entry["category"],
        path=entry["path"],
        row_count=len(rows),
    )

    if any(not header.strip() for header in headers):
        result.issues.append("빈 헤더가 있습니다.")

    for column_index, header in enumerate(headers, start=1):
        cell = worksheet.cell(1, column_index)
        cell.fill = WARNING_FILL if not header.strip() else HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    keys = entry["keys"]
    normalized_headers = {header.lower() for header in headers}
    missing_keys = [key for key in keys if key.lower() not in normalized_headers]
    if missing_keys:
        raise ValueError(f"{entry['sheet']} 필수 키 헤더 누락: {', '.join(missing_keys)}")

    seen_keys: dict[tuple[str, ...], int] = {}
    duplicate_rows: set[int] = set()
    for source_row_number, row in enumerate(rows, start=2):
        if len(row) != len(headers):
            result.issues.append(
                f"{source_row_number}행 열 개수 불일치: 헤더 {len(headers)}, 데이터 {len(row)}"
            )

        padded_row = row[: len(headers)] + [""] * max(0, len(headers) - len(row))
        converted_row: list[Any] = []
        cell_issues: list[tuple[int, str]] = []
        for column_index, (header, raw_value) in enumerate(
            zip(headers, padded_row), start=1
        ):
            converted, issue = convert_cell(header, raw_value)
            converted_row.append(converted)
            if issue:
                cell_issues.append((column_index, issue))

        worksheet.append(converted_row)
        for column_index, issue in cell_issues:
            cell = worksheet.cell(source_row_number, column_index)
            cell.fill = ERROR_FILL
            result.issues.append(
                f"{cell.coordinate}: {issue} 현재 값={cell.value!r}"
            )

        current_key = key_tuple(headers, padded_row, keys)
        if any(not value for value in current_key):
            result.issues.append(f"{source_row_number}행 필수 키가 비어 있습니다.")
        elif current_key in seen_keys:
            duplicate_rows.add(seen_keys[current_key])
            duplicate_rows.add(source_row_number)
            result.issues.append(
                f"{source_row_number}행 중복 키: {' / '.join(current_key)}"
            )
        else:
            seen_keys[current_key] = source_row_number

    for row_number in duplicate_rows:
        for cell in worksheet[row_number]:
            cell.fill = DUPLICATE_FILL

    max_row = max(worksheet.max_row, 2)
    max_column = max(worksheet.max_column, 1)
    worksheet.auto_filter.ref = f"A1:{get_column_letter(max_column)}{max_row}"
    worksheet.row_dimensions[1].height = 24
    add_data_validation(worksheet, headers, max_row)
    add_duplicate_rule(worksheet, headers, keys, max_row)
    set_column_widths(worksheet, headers, rows)
    return result


def build_index(workbook: Workbook, results: list[SheetResult]) -> None:
    worksheet = workbook["목차"]
    worksheet["A1"] = "NeoSurvive 밸런싱 데이터"
    worksheet["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    worksheet["A1"].fill = INDEX_TITLE_FILL
    worksheet.merge_cells("A1:F1")
    worksheet["A2"] = (
        "이 문서는 팀 편집용 원본입니다. Unity 적용 전 반드시 검증과 동기화를 거쳐야 합니다."
    )
    worksheet.merge_cells("A2:F2")
    worksheet["A2"].alignment = Alignment(wrap_text=True)

    headers = ["분류", "시트", "CSV 경로", "데이터 행", "상태", "검토 내용"]
    for column_index, header in enumerate(headers, start=1):
        cell = worksheet.cell(4, column_index, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for row_index, result in enumerate(results, start=5):
        worksheet.cell(row_index, 1, result.category)
        link_cell = worksheet.cell(row_index, 2, result.sheet)
        link_cell.hyperlink = f"#'{result.sheet}'!A1"
        link_cell.style = "Hyperlink"
        worksheet.cell(row_index, 3, result.path)
        worksheet.cell(row_index, 4, result.row_count)
        status_cell = worksheet.cell(row_index, 5, result.status)
        status_cell.fill = WARNING_FILL if result.issues else INDEX_SECTION_FILL
        issue_text = "\n".join(result.issues)
        issue_cell = worksheet.cell(row_index, 6, issue_text)
        issue_cell.alignment = Alignment(wrap_text=True, vertical="top")

    worksheet.freeze_panes = "A5"
    worksheet.auto_filter.ref = f"A4:F{4 + len(results)}"
    worksheet.column_dimensions["A"].width = 20
    worksheet.column_dimensions["B"].width = 24
    worksheet.column_dimensions["C"].width = 78
    worksheet.column_dimensions["D"].width = 12
    worksheet.column_dimensions["E"].width = 14
    worksheet.column_dimensions["F"].width = 72
    worksheet.row_dimensions[2].height = 36


def build_workbook(catalog: dict[str, Any], output_path: Path) -> list[SheetResult]:
    workbook = Workbook()
    index_sheet = workbook.active
    index_sheet.title = "목차"
    workbook.properties.title = catalog.get("workbook_title", "NeoSurvive Balancing")
    workbook.properties.subject = "Google Sheets 업로드용 밸런싱 데이터"

    results: list[SheetResult] = []
    for entry in catalog["entries"]:
        csv_path = PROJECT_ROOT / entry["path"]
        if not csv_path.is_file():
            raise FileNotFoundError(f"CSV 파일을 찾을 수 없습니다: {csv_path}")
        headers, rows = read_csv(csv_path)
        results.append(build_data_sheet(workbook, entry, headers, rows))

    build_index(workbook, results)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return results


def verify_workbook(catalog: dict[str, Any], output_path: Path) -> None:
    if not output_path.is_file():
        raise FileNotFoundError(f"검증할 워크북이 없습니다: {output_path}")

    workbook = load_workbook(output_path, read_only=False, data_only=False)
    expected_sheets = ["목차", *[entry["sheet"] for entry in catalog["entries"]]]
    if workbook.sheetnames != expected_sheets:
        raise ValueError(
            f"시트 구성이 다릅니다.\n예상: {expected_sheets}\n실제: {workbook.sheetnames}"
        )

    for entry in catalog["entries"]:
        worksheet = workbook[entry["sheet"]]
        csv_headers, csv_rows = read_csv(PROJECT_ROOT / entry["path"])
        workbook_headers = [
            worksheet.cell(1, column).value or ""
            for column in range(1, len(csv_headers) + 1)
        ]
        if workbook_headers != csv_headers:
            raise ValueError(f"{entry['sheet']} 헤더가 원본 CSV와 다릅니다.")
        if worksheet.max_row - 1 != len(csv_rows):
            raise ValueError(f"{entry['sheet']} 데이터 행 수가 원본 CSV와 다릅니다.")


def main() -> int:
    args = parse_args()
    try:
        catalog = load_catalog(args.catalog.resolve())
        output_path = args.output.resolve()

        if not args.verify_only:
            results = build_workbook(catalog, output_path)
            issue_count = sum(len(result.issues) for result in results)
            print(f"워크북 생성 완료: {output_path}")
            print(f"데이터 시트: {len(results)}개, 검토 항목: {issue_count}개")

        verify_workbook(catalog, output_path)
        print("워크북 구조 검증 완료: 목차 1개 + 데이터 시트 29개")
        return 0
    except Exception as error:
        print(f"오류: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
